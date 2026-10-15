import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for table cells in twips (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Set borders for the whole table."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def create_document():
    doc = Document()

    # Set page margins (Standard Vietnamese academic document: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # ~2.0 cm
        section.bottom_margin = Inches(0.79)  # ~2.0 cm
        section.left_margin = Inches(1.18)    # ~3.0 cm
        section.right_margin = Inches(0.79)   # ~2.0 cm

    # Set Normal Style Font
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(13)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    # 1. Header (2 columns: University & Republic)
    header_tbl = doc.add_table(rows=1, cols=2)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False
    
    col_widths = [Inches(3.0), Inches(3.4)]
    for i, col in enumerate(header_tbl.columns):
        col.width = col_widths[i]

    # Left cell: Trường & Khoa
    left_cell = header_tbl.cell(0, 0)
    set_cell_margins(left_cell, 0, 0, 0, 0)
    p_left1 = left_cell.paragraphs[0]
    p_left1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_left1.add_run("TRƯỜNG ĐẠI HỌC MỎ - ĐỊA CHẤT\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11.5)
    r1.font.bold = True
    
    r2 = p_left1.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11.5)
    r2.font.bold = True
    
    r_line1 = p_left1.add_run("-------------------")
    r_line1.font.name = 'Times New Roman'
    r_line1.font.size = Pt(10)

    # Right cell: Quốc hiệu
    right_cell = header_tbl.cell(0, 1)
    set_cell_margins(right_cell, 0, 0, 0, 0)
    p_right1 = right_cell.paragraphs[0]
    p_right1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p_right1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(11.5)
    r3.font.bold = True
    
    r4 = p_right1.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r4.font.name = 'Times New Roman'
    r4.font.size = Pt(12)
    r4.font.bold = True
    
    r_line2 = p_right1.add_run("-------------------------------")
    r_line2.font.name = 'Times New Roman'
    r_line2.font.size = Pt(10)

    doc.add_paragraph() # Spacing

    # 2. Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("ĐỀ CƯƠNG THỰC TẬP DOANH NGHIỆP")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    p_title.paragraph_format.space_after = Pt(12)

    # 3. Sections 1 to 5
    def add_section_header(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.bold = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        return p

    def add_info_line(label, value, bold_label=True):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run(label + ": ")
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(13)
        if bold_label:
            r_lbl.font.bold = True
        r_val = p.add_run(value)
        r_val.font.name = 'Times New Roman'
        r_val.font.size = Pt(13)
        return p

    # 1. Sinh viên thực tập
    add_section_header("1. Sinh viên thực tập")
    add_info_line("Họ và tên", "Nguyễn Ngọc Hữu")
    add_info_line("Mã sinh viên", "2221050720")
    add_info_line("Lớp", "DCCTCT67_08C")
    add_info_line("Điện thoại", "0344497635")
    add_info_line("Email", "nguyenhuu512gh@gmail.com")

    # 2. Giảng viên hướng dẫn
    add_section_header("2. Giảng viên hướng dẫn")
    add_info_line("Họ và tên", "Trần Trường Giang")
    add_info_line("Mã cán bộ", "(Đang cập nhật)")
    add_info_line("Điện thoại", "0948618185")
    add_info_line("Email", "trantruonggiang@humg.edu.vn")

    # 3. Thông tin đơn vị thực tập
    add_section_header("3. Thông tin đơn vị thực tập")
    add_info_line("Tên đơn vị", "Công ty cổ phần đầu tư AHV HOLDING")
    add_info_line("Địa chỉ", "Số nhà 18, Tổ 1, Phường Sông Công, Tỉnh Thái Nguyên, Việt Nam")
    add_info_line("Lĩnh vực hoạt động", "Lập trình website, các lĩnh vực công nghệ thông tin khác.")

    # 4. Người hướng dẫn thực tập tại công ty
    add_section_header("4. Người hướng dẫn thực tập tại công ty")
    add_info_line("Họ và tên", "Chu Văn Tú")
    add_info_line("Điện thoại", "0396261897")
    add_info_line("Email", "hotro@ahvholding.com")

    # 5. Vị trí thực tập
    add_section_header("5. Vị trí thực tập")
    p_pos = doc.add_paragraph()
    p_pos.paragraph_format.left_indent = Inches(0.2)
    p_pos.paragraph_format.space_after = Pt(4)
    r_pos = p_pos.add_run("Thực tập sinh AI-Driven Fullstack")
    r_pos.font.name = 'Times New Roman'
    r_pos.font.size = Pt(13)
    r_pos.font.bold = True

    # 6. Các nhiệm vụ được giao thực hiện (dự kiến)
    add_section_header("6. Các nhiệm vụ được giao thực hiện (dự kiến)")
    tasks = [
        ("Nghiên cứu & Khảo sát bài toán", "Khảo sát yêu cầu tra cứu tài liệu học vụ, quy chế nội bộ phục vụ dự án HUMG Ask. Nghiên cứu phương pháp lập trình AI-Driven Coding cùng quy trình xây dựng hệ thống hỏi đáp RAG (Retrieval-Augmented Generation) nâng cao với mô hình Google Gemini."),
        ("Thiết kế Cơ sở dữ liệu quan hệ & Vector", "Thiết kế kiến trúc CSDL trên nền tảng PostgreSQL (Neon Serverless) tích hợp phần mở rộng pgvector: quản lý tài khoản quản trị (admins), phiên làm việc (sessions), hội thoại (conversations), tin nhắn kèm trích dẫn (messages), tài liệu (documents) và phân đoạn vector 768 chiều (chunks)."),
        ("Xây dựng Pipeline RAG & Tìm kiếm lai (Hybrid Search)", "Xây dựng pipeline trích xuất văn bản đa định dạng (PDF, DOCX, TXT) với RecursiveCharacterTextSplitter; tạo vector ngữ nghĩa theo lô với Google Gemini Embeddings (gemini-embedding-001). Cài đặt giải thuật Hybrid Search kết hợp pgvector Cosine Distance và PostgreSQL Full-Text Search qua thuật toán tái xếp hạng RRF (Reciprocal Rank Fusion, k=60)."),
        ("Phát triển Backend RESTful API", "Xây dựng hệ thống Backend hoàn chỉnh trên nền tảng Node.js, Express.js 5 và TypeScript: bảo mật quản trị bằng JWT/Bcrypt, upload file tài liệu với Multer, kiểm định dữ liệu với Zod, quản lý phiên và tích hợp sinh câu trả lời thông minh bằng Google Gemini (gemini-3.1-flash-lite) kèm prompt chống hallucination."),
        ("Phát triển Frontend Web Application", "Xây dựng ứng dụng Single Page Application (SPA) bằng React 19, TypeScript, Vite và Tailwind CSS v4: trang Chatbot RAG với giao diện hiển thị câu trả lời Markdown trực quan và trích dẫn nguồn tài liệu tham khảo chi tiết; trang Admin Dashboard quản lý tài liệu, giám sát tiến trình vector hóa."),
        ("Kiểm thử & Đóng gói hệ thống", "Kiểm thử các luồng RESTful API bằng Postman; kiểm thử độ chính xác của phản hồi RAG; tối ưu hóa quy trình đóng gói và biên dịch sản phẩm; hoàn thiện quyển báo cáo và slide thuyết trình thực tập doanh nghiệp.")
    ]
    for title, desc in tasks:
        p_task = doc.add_paragraph()
        p_task.paragraph_format.left_indent = Inches(0.2)
        p_task.paragraph_format.space_after = Pt(4)
        r_bullet = p_task.add_run("• ")
        r_bullet.font.bold = True
        r_t = p_task.add_run(f"{title}: ")
        r_t.font.bold = True
        r_d = p_task.add_run(desc)
        r_bullet.font.name = r_t.font.name = r_d.font.name = 'Times New Roman'
        r_bullet.font.size = r_t.font.size = r_d.font.size = Pt(13)

    # 7. Kế hoạch thực hiện (dự kiến)
    add_section_header("7. Kế hoạch thực hiện (dự kiến)")

    # Table 10 weeks
    weeks_data = [
        ("1", 
         "• Tiếp nhận vị trí thực tập tại AHV HOLDING, làm quen văn hóa doanh nghiệp và quy định bảo mật thông tin.\n• Tiếp nhận và thiết lập môi trường lập trình hỗ trợ AI-Driven Coding (Node.js, TypeScript, Vite, Git, AI Coding Assistant).\n• Khảo sát bài toán tra cứu văn bản quy chế, tài liệu học vụ nội bộ (dự án HUMG Ask).",
         "• Báo cáo tổng quan doanh nghiệp và bài toán nghiệp vụ.\n• Môi trường lập trình (IDE, Git repo, Node.js 20+) đã hoàn tất và cấu hình chuẩn."),
        
        ("2",
         "• Thu thập, phân loại và chuẩn hóa tập văn bản tài liệu mẫu (các định dạng PDF, DOCX, TXT).\n• Nghiên cứu kiến trúc RAG (Retrieval-Augmented Generation) với thư viện LangChain (LangChain.js) và hệ sinh thái mô hình Google Gemini (Gemini Embeddings, Gemini Flash).\n• Nghiên cứu kỹ thuật tìm kiếm lai (Hybrid Search: Vector Search kết hợp Full-Text Search).",
         "• Tập dữ liệu tài liệu mẫu đã chuẩn hóa.\n• Báo cáo lý thuyết về kiến trúc RAG, LangChain và giải pháp giải thuật RRF đề xuất."),
        
        ("3",
         "• Thiết kế kiến trúc tổng thể hệ thống Fullstack (Client SPA - Express Backend - PostgreSQL pgvector - Google Gemini AI).\n• Thiết kế CSDL quan hệ kết hợp Vector Database trên PostgreSQL (Neon Serverless): bảng admins, sessions, conversations, messages, documents và bảng chunks (trường embedding kiểu vector(768)).",
         "• Sơ đồ kiến trúc hệ thống, sơ đồ CSDL (ERD quan hệ & vector).\n• Script khởi tạo schema CSDL PostgreSQL tích hợp extension pgvector.\n• Báo cáo tiến độ tuần 1-3 có xác nhận của GVHD."),
        
        ("4",
         "• Xây dựng Document Pipeline: trích xuất nội dung từ PDF (pdf-parse), DOCX (mammoth), TXT; phân đoạn văn bản bằng RecursiveCharacterTextSplitter.\n• Xây dựng module sinh embedding theo lô với Google Gemini Embeddings (gemini-embedding-001, 768 chiều) và lưu trữ an toàn trong Database Transaction (BEGIN...COMMIT/ROLLBACK).\n• Cài đặt giải thuật Hybrid Search kết hợp giữa pgvector Cosine Distance và PostgreSQL Full-Text Search qua thuật toán RRF (Reciprocal Rank Fusion, k=60).",
         "• Module xử lý tài liệu, chunking và sinh vector hoạt động chính xác.\n• Hàm truy vấn tìm kiếm lai Hybrid Search + RRF trích xuất đúng top 5 chunks ngữ cảnh liên quan."),
        
        ("5",
         "• Phát triển Backend API bằng Node.js, Express.js 5 và TypeScript.\n• Xây dựng hệ thống xác thực bảo mật Quản trị viên (Admin Login) bằng JWT (jsonwebtoken) và mã hóa mật khẩu bcryptjs.\n• Xây dựng các Endpoints quản lý tài liệu: upload file đa định dạng qua Multer, kiểm tra dữ liệu đầu vào bằng Zod, xóa tài liệu và tải file gốc.\n• Áp dụng phương pháp AI-Driven Coding để tối ưu hóa mã nguồn.",
         "• Bộ RESTful APIs quản trị tài liệu và xác thực hoạt động ổn định (/admin/login, /documents/create, /documents/list, /documents/:id, /documents/:id/file).\n• Seed script khởi tạo tài khoản quản trị đầu tiên.\n• Tài liệu API Documentation (Postman Collection)."),
        
        ("6",
         "• Tiếp tục phát triển Backend API: Xây dựng các Endpoints quản lý phiên làm việc (/session/create), tạo và quản lý cuộc trò chuyện (/conversation).\n• Xây dựng Endpoint hỏi đáp RAG (POST /conversation/messages): tích hợp lịch sử trò chuyện gần nhất, context trích xuất từ Hybrid Search, kiểm soát prompt chống ảo giác và gọi Google Gemini LLM (gemini-3.1-flash-lite).\n• Trả về câu trả lời kèm danh sách nguồn trích dẫn (sources: documentId, title, page, similarity).\n• Nộp Báo cáo Tiến độ 2 (Phụ lục 2).",
         "• Tối thiểu 08 RESTful APIs hoàn chỉnh (đầy đủ GET, POST, DELETE).\n• Chatbot RAG Backend trả lời đúng ngữ cảnh tài liệu và trích dẫn chính xác số trang/nguồn.\n• Báo cáo tiến độ tuần 4-6 có xác nhận của GVHD."),
        
        ("7",
         "• Thiết kế UI/UX và lập trình Frontend Single Page Application bằng React 19, TypeScript, Vite và Tailwind CSS v4.\n• Xây dựng trang giao diện Chatbot RAG (ChatPage), Sidebar lịch sử cuộc trò chuyện, khung nhập liệu tin nhắn.\n• Tích hợp thư viện react-markdown và remark-gfm để render câu trả lời chatbot chuẩn format Markdown, hiển thị khối nguồn tài liệu tham khảo trực quan.",
         "• Giao diện trang Chatbot RAG hoàn chỉnh, hỗ trợ tương tác mượt mà, thân thiện trên các thiết bị.\n• Kết nối API chat với Backend thành công."),
        
        ("8",
         "• Phát triển các trang giao diện Quản trị: Trang quản lý tài liệu (AdminDashboardPage), Modal tải tài liệu lên hệ thống (AdminUploadModal), xem trạng thái xử lý (processing/active/failed).\n• Phát triển Trang đăng nhập quản trị viên (LoginPage), xây dựng cơ chế bảo vệ định tuyến (ProtectedRoute).\n• Tích hợp toàn diện Fullstack giữa Frontend (React/Vite) và Backend (Express API).",
         "• Hệ thống Web Fullstack hoàn chỉnh 03 màn hình chính (ChatPage, AdminDashboardPage, LoginPage) hoạt động đồng bộ.\n• Luồng nghiệp vụ từ Upload tài liệu -> Vector hóa -> Chat hỏi đáp diễn ra trơn tru."),
        
        ("9",
         "• Kiểm thử toàn bộ hệ thống: Kiểm thử API bằng Postman, kiểm thử giao diện người dùng và đo lường độ chính xác câu trả lời RAG.\n• Tối ưu hóa hiệu năng xây dựng ứng dụng (Vite Production Build) và rà soát cấu hình biến môi trường (.env).\n• Đánh giá khả năng xử lý lỗi (file scan không text, lỗi quá giới hạn trang, xử lý rollback dữ liệu).\n• Xin xác nhận và đóng dấu Công lệnh tại doanh nghiệp.",
         "• Bảng kết quả kiểm thử hệ thống (Testcases & Postman collection).\n• Bản build sản phẩm hoàn thiện sẵn sàng triển khai.\n• Công lệnh thực tập có đóng dấu xác nhận của Công ty cổ phần đầu tư AHV HOLDING."),
        
        ("10",
         "• Hoàn thiện Quyển Báo cáo Thực tập doanh nghiệp (theo mẫu Phụ lục 3), Bảng khai báo sử dụng công cụ AI trong quá trình thực hiện.\n• Xây dựng Slide thuyết trình báo cáo TTDN theo chuẩn quy định của Khoa CNTT.\n• Giảng viên hướng dẫn kiểm tra sản phẩm hoàn thiện và duyệt cho bảo vệ.",
         "• Quyển Báo cáo TTDN hoàn chỉnh.\n• Slide thuyết trình TTDN.\n• Mã nguồn hoàn chỉnh của dự án trên Git repository.")
    ]

    tbl_plan = doc.add_table(rows=1, cols=3)
    tbl_plan.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_plan.autofit = False
    set_table_borders(tbl_plan, color="999999", sz="4", val="single")

    col_widths = [Inches(0.8), Inches(3.2), Inches(2.5)]
    for i, col in enumerate(tbl_plan.columns):
        col.width = col_widths[i]

    # Header Row
    hdr_cells = tbl_plan.rows[0].cells
    headers = ["Tuần", "Nội dung công việc làm theo tuần", "Kết quả dự kiến"]
    for i, header_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_margins(hdr_cells[i], 120, 120, 140, 140)
        set_cell_shading(hdr_cells[i], "E6EDF5")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(header_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True

    # Data Rows
    for w_num, w_task, w_res in weeks_data:
        row_cells = tbl_plan.add_row().cells
        for idx in range(3):
            row_cells[idx].width = col_widths[idx]
            set_cell_margins(row_cells[idx], 100, 100, 120, 120)
        
        # Col 0: Tuần
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(w_num)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.font.bold = True

        # Col 1: Nội dung công việc
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(w_task)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11.5)

        # Col 2: Kết quả dự kiến
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(w_res)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11.5)

    doc.add_paragraph() # Spacing

    # Date and Signatures
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_date = p_date.add_run("Hà Nội, ngày       tháng       năm 2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.font.italic = True
    p_date.paragraph_format.space_after = Pt(10)

    # 3-column Signatures Table
    sig_tbl = doc.add_table(rows=2, cols=3)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    
    sig_col_widths = [Inches(2.2), Inches(2.2), Inches(2.2)]
    for i, col in enumerate(sig_tbl.columns):
        col.width = sig_col_widths[i]

    roles = [
        ("Người hướng dẫn tại công ty", "(Ký, ghi rõ họ tên)", "Chu Văn Tú"),
        ("Giảng viên hướng dẫn", "(Ký, ghi rõ họ tên)", "Trần Trường Giang"),
        ("Sinh viên thực tập", "(Ký, ghi rõ họ tên)", "Nguyễn Ngọc Hữu")
    ]

    for idx, (role, sub, name) in enumerate(roles):
        cell_top = sig_tbl.cell(0, idx)
        set_cell_margins(cell_top, 0, 0, 0, 0)
        p_r = cell_top.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_r.paragraph_format.space_after = Pt(2)
        r_role = p_r.add_run(role + "\n")
        r_role.font.name = 'Times New Roman'
        r_role.font.size = Pt(12)
        r_role.font.bold = True

        r_sub = p_r.add_run(sub)
        r_sub.font.name = 'Times New Roman'
        r_sub.font.size = Pt(11)
        r_sub.font.italic = True

        # Bottom cell with spacing for real signature
        cell_bot = sig_tbl.cell(1, idx)
        set_cell_margins(cell_bot, 0, 0, 0, 0)
        p_n = cell_bot.paragraphs[0]
        p_n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_n.paragraph_format.space_before = Pt(65) # Space for signature
        p_n.paragraph_format.space_after = Pt(0)
        r_name = p_n.add_run(name)
        r_name.font.name = 'Times New Roman'
        r_name.font.size = Pt(12.5)
        r_name.font.bold = True

    # Page Break for Technology Stack Annex
    doc.add_page_break()

    # Annex Title
    p_annex_title = doc.add_paragraph()
    p_annex_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_annex_title = p_annex_title.add_run("PHỤ LỤC: BẢNG TỔNG HỢP CÔNG NGHỆ THỰC TẾ DỰ ÁN HUMG ASK V3")
    r_annex_title.font.name = 'Times New Roman'
    r_annex_title.font.size = Pt(14)
    r_annex_title.font.bold = True
    p_annex_title.paragraph_format.space_after = Pt(12)

    tech_stack = [
        ("Kiến trúc Backend", "Node.js (v20+), Express.js 5.2.1, TypeScript 5.9, ts-node, nodemon."),
        ("Kiến trúc Frontend", "React 19, Vite 8, TypeScript, Tailwind CSS v4, Lucide React (Icons), React Router DOM v7, React Markdown & Remark GFM."),
        ("Cơ sở dữ liệu & Vector DB", "PostgreSQL (Neon Serverless) tích hợp tiện ích mở rộng pgvector. Lưu trữ embeddings vector 768 chiều trực tiếp trong CSDL với Transaction ACID an toàn."),
        ("Khung làm việc AI & RAG", "LangChain.js (@langchain/core, @langchain/google-genai, @langchain/textsplitters)."),
        ("Mô hình AI (Google Gemini)", "• Embedding Model: gemini-embedding-001 (768 dimensions, TaskType: RETRIEVAL_DOCUMENT / RETRIEVAL_QUERY).\n• Chat LLM: gemini-3.1-flash-lite (temperature: 0.2, prompt chống ảo giác)."),
        ("Giải thuật tìm kiếm", "Hybrid Search kết hợp Vector Cosine Distance (<=>) và PostgreSQL Full-Text Search (ts_rank_cd) thông qua giải thuật Reciprocal Rank Fusion (RRF, k=60) trích xuất top 5 chunk chuẩn xác nhất."),
        ("Xử lý & Phân tách tài liệu", "Hỗ trợ đa định dạng: PDF (pdf-parse), DOCX (mammoth), TXT (TextDecoder UTF-8). Phân đoạn văn bản bằng RecursiveCharacterTextSplitter (chunk size: 1000, overlap: 120 ký tự)."),
        ("Bảo mật & Quản lý phiên", "JSON Web Token (jsonwebtoken), Bcrypt (bcryptjs), Cookie Session quản lý khách vãng lai, Zod schema validation dữ liệu đầu vào.")
    ]

    tbl_tech = doc.add_table(rows=1, cols=2)
    tbl_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tech.autofit = False
    set_table_borders(tbl_tech, color="999999", sz="4", val="single")

    t_widths = [Inches(2.2), Inches(4.3)]
    for i, col in enumerate(tbl_tech.columns):
        col.width = t_widths[i]

    # Header
    h0 = tbl_tech.cell(0, 0)
    h1 = tbl_tech.cell(0, 1)
    set_cell_margins(h0, 100, 100, 120, 120)
    set_cell_margins(h1, 100, 100, 120, 120)
    set_cell_shading(h0, "E6EDF5")
    set_cell_shading(h1, "E6EDF5")
    
    p = h0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Thành phần / Hạng mục")
    r.font.bold = True
    r.font.size = Pt(12)

    p = h1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Công nghệ & Thư viện sử dụng trong dự án")
    r.font.bold = True
    r.font.size = Pt(12)

    for item, detail in tech_stack:
        row = tbl_tech.add_row().cells
        row[0].width = t_widths[0]
        row[1].width = t_widths[1]
        set_cell_margins(row[0], 90, 90, 110, 110)
        set_cell_margins(row[1], 90, 90, 110, 110)

        p_it = row[0].paragraphs[0]
        p_it.paragraph_format.space_after = Pt(0)
        r_it = p_it.add_run(item)
        r_it.font.name = 'Times New Roman'
        r_it.font.size = Pt(11.5)
        r_it.font.bold = True

        p_dt = row[1].paragraphs[0]
        p_dt.paragraph_format.space_after = Pt(0)
        p_dt.paragraph_format.line_spacing = 1.15
        r_dt = p_dt.add_run(detail)
        r_dt.font.name = 'Times New Roman'
        r_dt.font.size = Pt(11.5)

    # Save to file
    output_filename = "DE_CUONG_THUC_TAP_DOANH_NGHIEP_NGUYEN_NGOC_HUU.docx"
    doc.save(output_filename)
    print(f"Document created successfully: {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    create_document()
